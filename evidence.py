import os
import re
import xml.etree.ElementTree as ET
from html import unescape
from docx import Document
from docx.shared import Inches, RGBColor

STATUS_COLOR = {
    'PASS': RGBColor(0, 128, 0),
    'FAIL': RGBColor(255, 0, 0),
    'NOT RUN': RGBColor(128, 128, 128),
}

def extrair_fluxo(element, test_name):
    resultado = []
    for kw in element.findall('kw'):
        name = unescape(kw.attrib.get('name', ''))
        owner = kw.attrib.get('owner', '')
        status_el = kw.find('status')
        status = status_el.attrib.get('status', 'PASS') if status_el is not None else 'PASS'
        msg_el = kw.find('msg')
        time = msg_el.attrib.get('time') if msg_el is not None else '9999-99-99T99:99:99.999'
        args = [unescape(arg.text or '') for arg in kw.findall('arg')]
        argumentos_formatados = f" {args}" if args else ""

        entrada = {
            'test_name': test_name,
            'name': name + argumentos_formatados,
            'owner': owner,
            'status': status,
            'time': time,
            'screenshot': None,
        }

        for subkw in kw.findall('.//kw'):
            if subkw.attrib.get('owner') == 'Browser' and subkw.attrib.get('name') == 'Take Screenshot':
                msg_sub = subkw.find('msg')
                if msg_sub is not None:
                    match = re.search(r'href="(.*?)"', msg_sub.text or '')
                    if match:
                        entrada['screenshot'] = unescape(match.group(1))

        resultado.append(entrada)
        resultado.extend(extrair_fluxo(kw, test_name))
    return resultado

def gerar_docx(output_xml_path, pasta_raiz_screenshot, docx_saida, template_path=None):
    tree = ET.parse(output_xml_path)
    root = tree.getroot()
    doc = Document(template_path) if template_path and os.path.exists(template_path) else Document()

    todas_entradas = []
    for test in root.findall('.//test'):
        test_name = test.attrib.get('name', 'Sem nome')
        todas_entradas.extend(extrair_fluxo(test, test_name))

    todas_entradas.sort(key=lambda x: x['time'])

    ultimo_teste = None
    for entrada in todas_entradas:
        # if entrada['owner'] != 'keywords':
        if entrada['owner'] != 'keywords' or entrada['name'].strip().startswith("Gerar evidencia"):


            continue

        if entrada['test_name'] != ultimo_teste:
            doc.add_page_break() if ultimo_teste else None
            doc.add_heading(f'Teste: {entrada["test_name"]}', level=1)
            ultimo_teste = entrada['test_name']

        texto = f'FALHOU: {entrada["name"]}' if entrada['status'] == 'FAIL'  else entrada['name']
        p = doc.add_paragraph(texto)
        run = p.runs[0]
        run.font.color.rgb = STATUS_COLOR.get(entrada['status'], RGBColor(0, 0, 0))

        if entrada['screenshot']:
            caminho_img = os.path.join(pasta_raiz_screenshot, os.path.basename(entrada['screenshot']))
            if os.path.exists(caminho_img):
                doc.add_picture(caminho_img, width=Inches(5.5))
            else:
                doc.add_paragraph(f"[Imagem não encontrada: {caminho_img}]")

    doc.save(docx_saida)
    print(f"📄 Documento '{docx_saida}' gerado com sucesso.")

# === EXECUTAR ===
if __name__ == "__main__":
    gerar_docx(
        output_xml_path="output.xml",
        pasta_raiz_screenshot="browser/screenshot",
        docx_saida="evidencia_final.docx",
        template_path="template_evidencia.docx"  # ou None
    )
